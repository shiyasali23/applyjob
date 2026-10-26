"""Reports generated from the ledger: tracker spreadsheet, status summary, audit findings sheet, LinkedIn keyword gaps."""
import collections, json
from pathlib import Path
from . import ledger, vocab, workspace, profile as P

ORDER = ["SUBMITTED", "EMAILED", "APPROVED", "READY", "NEEDS_USER", "LOGIN", "EXPIRED", "EMAIL_ONLY", "AUDIT_FAIL",
         "AUDITED", "DRAFTED", "GOOD", "CHECK", "FOUND", "INTERVIEW", "OFFER", "REPLIED", "REJECTED", "CLOSED", "SKIP", "REMOVED"]


def tracker(db, out=None):
    import openpyxl
    from openpyxl.styles import Font
    out = Path(out or workspace.path("tracker.xlsx"))
    rows = ledger.export_rows(db)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Applications"
    head = ["id", "company", "title", "location", "status", "sent", "evidence", "note", "batch", "folder", "url", "source", "updated"]
    ws.append(head)
    for c in ws[1]: c.font = Font(bold=True)
    for r in rows: ws.append([r[h] for h in head])
    ws.freeze_panes = "A2"
    tmp = out.with_suffix(".tmp.xlsx"); wb.save(tmp); tmp.replace(out)
    return out, len(rows)


def summary(db, batch=None):
    c = ledger.counts(db, batch)
    lines = ["%s: %d jobs | sent %d" % (batch or "all batches", sum(c.values()), c.get("SUBMITTED", 0) + c.get("EMAILED", 0))]
    for st in ORDER:
        if c.get(st):
            js = ledger.jobs(db, st, batch)
            lines.append("\n%s (%d)" % (st, len(js)))
            for j in js[:200]:
                lines.append("  %4d  %s | %s%s" % (j["id"], j["company"][:40], j["title"][:50], ("  -> " + j["note"][:90]) if j["note"] and st not in ("SUBMITTED", "EMAILED") else ""))
    return "\n".join(lines)


def findings_sheet(batch_dir, out=None):
    """Every audit finding of a batch, one row each (the atomic error list)."""
    import openpyxl
    B = Path(batch_dir); out = Path(out or B / "_apply" / "findings.xlsx"); out.parent.mkdir(exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(["job", "id", "severity", "where", "rule", "evidence", "fix"])
    n = collections.Counter()
    for f in sorted(B.glob("*/audit.json")):
        for x in json.loads(f.read_text()):
            ws.append([f.parent.name, x["id"], x["severity"], x["where"], x["rule"], x["evidence"], x["fix"]]); n[x["severity"]] += 1
    wb.save(out)
    return out, dict(n)


def linkedin_gaps(db, prof, profile_text, top=40):
    """Skills the target adverts ask for most, and whether the LinkedIn profile text contains them."""
    cnt = collections.Counter()
    for j in ledger.jobs(db, ["GOOD", "DRAFTED", "AUDITED", "READY", "APPROVED", "SUBMITTED", "EMAILED", "NEEDS_USER"]):
        for t in set(vocab.find_terms(j["description"] or "")): cnt[t] += 1
    known = P.known_skills(prof)
    rows = []
    for t, n in cnt.most_common(top):
        rows.append({"term": t, "adverts": n, "on_linkedin": bool(vocab.rx(t).search(profile_text or "")),
                     "in_profile": P.norm(t) in known})
    return rows
